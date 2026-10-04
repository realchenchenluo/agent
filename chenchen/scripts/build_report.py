"""Build the Chinese product review. Requires python-docx; not required to run the apps."""
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.opc.constants import RELATIONSHIP_TYPE as RT

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / 'docs' / '两版产品迭代与店主持续使用设计.docx'
doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Inches(8.5), Inches(11)
sec.top_margin, sec.bottom_margin = Inches(.8), Inches(.75)
sec.left_margin = sec.right_margin = Inches(1)
sec.header_distance = sec.footer_distance = Inches(.35)
for name, size in [('Normal',11),('Title',26),('Subtitle',12),('Heading 1',18),('Heading 2',13)]:
    style = doc.styles[name]
    style.font.name = 'Microsoft YaHei'
    style.font.size = Pt(size)
    style.font.color.rgb = RGBColor(0,0,0)
    style.font.italic = False
    style.font.bold = name.startswith('Heading')
    fonts=style.element.get_or_add_rPr().rFonts
    for attr in ['asciiTheme','hAnsiTheme','eastAsiaTheme','cstheme']:
        fonts.attrib.pop(qn('w:'+attr),None)
    for attr in ['ascii','hAnsi','eastAsia','cs']:
        fonts.set(qn('w:'+attr),'Microsoft YaHei')
    style.paragraph_format.space_after = Pt(6)
    style.paragraph_format.line_spacing = Pt(15.5) if name == 'Normal' else 1.1
    snap=OxmlElement('w:snapToGrid');snap.set(qn('w:val'),'0');style.element.get_or_add_pPr().append(snap)
    style.paragraph_format.widow_control = True
for name in ['Heading 1','Heading 2']:
    doc.styles[name].paragraph_format.space_before = Pt(12)
    doc.styles[name].paragraph_format.keep_with_next = True
doc.core_properties.title = '两版产品迭代与店主持续使用设计'
doc.core_properties.subject = 'Portfolio Sentinel v0.4.0 与店主财务助手'
doc.core_properties.author = 'chenchen'
doc.core_properties.keywords = '产品评审,账单,现金流,Agent,演示版本'
hp=sec.header.paragraphs[0]
hp.text='CHENCHEN  /  PRODUCT REVIEW  /  2026.10.04'
hp.runs[0].font.size=Pt(8)
hp.runs[0].font.color.rgb=RGBColor.from_string('666666')
fp=sec.footer.paragraphs[0]
fp.alignment=WD_ALIGN_PARAGRAPH.RIGHT
fp.add_run('两版迭代评审 · 演示版本     ').font.size=Pt(8)
field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');fp._p.append(field)

def p(text, style=None):
    return doc.add_paragraph(text,style)

def h(text, level=2):
    return doc.add_heading(text,level)

def page(title):
    heading=h(title,1)
    heading.paragraph_format.page_break_before=True

def table(headers, rows, widths):
    t=doc.add_table(rows=1, cols=len(headers))
    t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
    for c,w in zip(t.columns,widths):c.width=Inches(w)
    for cell,text,w in zip(t.rows[0].cells,headers,widths):cell.text=text;cell.width=Inches(w)
    for row in rows:
        cells=t.add_row().cells
        for cell,text,w in zip(cells,row,widths):cell.text=str(text);cell.width=Inches(w)
    for i,row in enumerate(t.rows):
        trpr=row._tr.get_or_add_trPr()
        trpr.append(OxmlElement('w:cantSplit'))
        if i==0:trpr.append(OxmlElement('w:tblHeader'))
        for cell in row.cells:
            cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tcpr=cell._tc.get_or_add_tcPr()
            margins=OxmlElement('w:tcMar')
            for side in ['top','left','bottom','right']:
                e=OxmlElement('w:'+side);e.set(qn('w:w'),'80');e.set(qn('w:type'),'dxa');margins.append(e)
            tcpr.append(margins)
            if i==0:
                shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'EDEFEA');tcpr.append(shade)
            for para in cell.paragraphs:
                para.paragraph_format.space_after=Pt(2)
                para.paragraph_format.space_before=Pt(2)
                para.paragraph_format.line_spacing=Pt(14)
                for run in para.runs:
                    run.font.size=Pt(10)
                    run.bold=i==0
    borders=OxmlElement('w:tblBorders')
    for side in ['top','left','bottom','right','insideH','insideV']:
        e=OxmlElement('w:'+side);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'CDD2C9');borders.append(e)
    t._tbl.tblPr.append(borders)
    gap=p('');gap.paragraph_format.line_spacing=Pt(1);gap.paragraph_format.space_after=Pt(2)
    return t

def link(label, url):
    para=p('')
    rel=para.part.relate_to(url,RT.HYPERLINK,is_external=True)
    hyp=OxmlElement('w:hyperlink');hyp.set(qn('r:id'),rel)
    run=OxmlElement('w:r');props=OxmlElement('w:rPr')
    color=OxmlElement('w:color');color.set(qn('w:val'),'245B49');props.append(color)
    size=OxmlElement('w:sz');size.set(qn('w:val'),'20');props.append(size)
    run.append(props);text=OxmlElement('w:t');text.text=label;run.append(text);hyp.append(run);para._p.append(hyp)
    return para

p('两版产品迭代与\n店主持续使用设计','Title')
p('持仓健康管理优化版 × 店主财务助手新增版','Subtitle')
p('版本 v0.4.0  ·  评审日期 2026 年 10 月 4 日')
h('这次交付的核心结果')
p('保留“让每一个人，都拥有自己的基金经理”的产品方向，本次交付拆成两个独立产品。它们共用本地服务的运行能力，但拥有各自的品牌、首屏重点和使用路径：投资版先看风险与改进方案；店主版先看余额与账单，再进入现金安排。')
table(['版本','服务的问题','本次可运行结果'],[
    ['01 持仓健康管理','组合是否集中？几个调整方案有什么差别？','风险偏好收集、持仓校验、结构诊断、A/B/C 同口径模拟、证据导出。'],
    ['02 店主财务助手','钱何时到账？支出为何变化？接下来哪天付款？','多渠道账本、净到账拆解、14 天余额情景、采购解释、核对后提醒与关注。']
],[1.15,2.35,3.0])
h('怎样判断功能值得做')
p('投资版继续追问：“一个真正属于我的基金经理，会不会做这件事？”店主版对应的问题是：“一个了解我的店、尊重我的判断的财务助手，会不会先解释这笔账，再提醒我行动？”')
p('因此，本次优先完成可追溯计算、明确的日期口径、人工确认和可继续查看的记录；暂不引入真实支付、自动交易或没有凭证支持的强结论。')
h('阅读方式')
p('第 2 页看投资版改动；第 3—4 页看店主账单与使用流程；第 5 页看持续使用设计；第 6 页看修改过程与职责边界；第 7 页看验证与限制；第 8 页看运行和资料来源。页面设计按“首屏突出重点、点击进入详情”组织。')
p('所有金额与证券均为虚构演示。下文的用户留存目标是待验证假设，不是已经获得的用户研究结果。')

page('01  投资版：让比较结果可以复算')
p('新入口由后台真实计算生成诊断与方案表，原始静态素材保留用于对照。界面改成独立的投资顾问产品：首屏先展示风险结果，A/B/C 方案以卡片突出，曲线、目标权重和运行证据收进可展开详情。分析流程仍停在团队风控交接点，不冒充已完成审批或真实成交。')
table(['改动','实现方式','用户得到什么'],[
    ['输入更严格','数值、份额、固定日期与价格逐项核对；带 snapshot_id 的输入也复核金额和权重。','不把字符串、伪造权重或错期价格当成有效持仓。'],
    ['同向风险识别','只将相关系数 ≥0.8 的资产对列为同向集中；现金占比独立展示。','负相关不再被误列成同向聚集。'],
    ['统一费用口径','净值曲线与收益指标用同一费用后路径；A 保持方案换手为零。','表格收益与曲线终点一致。'],
    ['情景可比','联合资产收益连续区块重采样；四组使用相同 seed 与抽样索引。','减少方案差异来自随机抽样的干扰。'],
    ['新资产参与计算','价格路径遍历目标资产集，而非仅遍历原有持仓。','方案加入债券 ETF 时不会漏算它。'],
    ['运行证据与恢复','保留 ToolResult 原文；输入改变清除旧 checkpoint；同输入复用；未知任务报错。','能追溯来源，避免拿旧结果解释新输入。']
],[1.0,2.85,2.65])
h('固定演示可以核对的结果')
p('组合总额 31,080 元，现金 10,000 元（32.18%），最大证券仓位为 ETF_GROWTH（24.29%）。短历史窗口回撤约 −2.79%。Current 与 A 的收益、换手和情景分位数一致；B 换手约 4.29%，C 约 12.79%。')
h('研究启发与适用边界')
p('参考 scikit-learn 对数据泄漏的说明 [3]，明确把当前模拟标为“事后配置比较”。只有 7 个虚构价格观测，不能验证策略在未知市场中的效果。年化波动、2,000 条路径与分位数仅用于展示计算方法；增加路径数不能补足原始样本。')
p('风险偏好进入任务状态并随交接传递；A/B/C 仍按固定规则生成，尚不构成个性化适配。超出仓位或换手约束的输入会中止，不能悄悄放宽规则。')

page('02  店主版：先统一钱的口径')
p('示例为“巷口咖啡”，基准日 2026-10-04，共 15 笔虚构账单。平台名称只用来展示分散渠道，不代表已对接这些平台，也不代表真实费率或结算周期。')
table(['账户视图','金额','解释'],[
    ['已核对余额','6,200.00 元','由店主手动核对；已结算收支已经包含在内。'],
    ['还在路上的钱','7,176.80 元','包含 855 元逾期待收款；尚不是可用余额。'],
    ['已知待付款','16,150.00 元','来自已录入的供应商、房租、工资和水电账单。'],
    ['外卖结算批次','2,744.00 元净收','原额 3,200 − 平台费 256 − 退款 200。']
],[1.4,1.2,3.9])
h('14 天现金表：日期比一个总数更有用')
p('预计余额 = 前日余额 + 当日预计净到账 − 当日应付款。逾期收入不默认到账；逾期应付款预留在次日。另一条保守线只扣已知付款，假设待到账款暂时全部未收到。')
table(['日期','预计到账','到期付款','预计期末余额'],[
    ['10 月 5 日','3,577.80','2,200.00','7,577.80'],
    ['10 月 6 日','2,744.00','7,000.00','3,321.80'],
    ['10 月 7 日','0.00','1,100.00','2,221.80'],
    ['10 月 8 日','0.00','400.00','1,821.80'],
    ['10 月 9 日','0.00','650.00','1,171.80'],
    ['10 月 10 日','0.00','4,800.00','−3,628.20']
],[1.5,1.5,1.5,2.0])
p('表内单位：元。按预计日期收款，首次缺口在 10 月 10 日；收款整体延迟 3 天，首次缺口提前到 10 月 6 日，当日余额为 −3,000 元。付款日期不会跟着收款延迟而自动后移。')
p('参考 Stripe 对结算批次与交易明细的对应关系 [1]，保留原额、扣费、退款和净额；参考 Xero 的已知账单现金预测思路 [2]，只测算已知收支。两者都是产品机制参考，不是中国支付平台规则的依据。')

page('03  店主版：先解释，再提醒')
h('一条可以亲自体验的路径')
table(['步骤','页面行为','系统约束'],[
    ['1 看总览','首屏突出当前余额、待到账、待付款和首个缺口。','已结算款不会再被加到当前余额；完整日历收进详情。'],
    ['2 看原始账单','首屏先给三张账单卡；需要时展开筛选和完整账单，点击“看账单”。','明细显示来源引用与原额、扣费、退款、净额。'],
    ['3 确认口径','明确勾选理解余额和预计日期的区别。','服务端要求已查看至少一笔来源、confirmed=true、当前 revision。'],
    ['4 看提醒','生成现金缺口、逾期待收和采购变化提醒。','每条提醒可回到相关账单；没有证据不解释涨价原因。'],
    ['5 下次继续','标记关注；导出账单或核对摘要。','刷新可恢复会话；导入新账单清除旧确认与关注。']
],[1.1,2.4,3.0])
h('采购量变多，不等于东西涨价')
p('按记账日比较两个等长 7 天窗口：本期 9 月 28 日至 10 月 4 日，上期 9 月 21 日至 27 日。同一种咖啡豆从 20 kg、1,600 元，变为 25 kg、2,200 元；采购支出增加 600 元（37.5%）。')
p('数量变化影响 = (25 − 20) × 80 = 400 元。单价变化影响 = (88 − 80) × 25 = 200 元。两项合计恰为实际支出增量。单价、品名和单位都能回到账单核对。')
p('如果品名不同、单位不同或缺少上期记录，只展示支出变化，不断言“供应商涨价”。采购支出也不等于当期实际耗用成本，不能直接推出毛利下降。')
h('导入时保护已核对的账')
p('只支持规范 JSON、人民币与整数分。以“渠道 + 账单号”识别重复；完全相同的账单去重，金额或状态冲突则整批拒绝，不局部覆盖。导入前先验证，失败时原有账本和确认状态保持不变。成功替换后必须重新查看与确认。')
p('本轮“账单来源”是来源字符串，不是上传的凭证图片或自动验证的银行流水；看过一笔来源也不意味着所有凭证已经审计。')

page('04  怎样成为店主愿意持续使用的工具')
p('店主回来的理由应是一件经营事务有了变化：一笔款到了、一张单该付了、一个价格变了。持续使用的价值来自少一次遗漏、少一次翻找和更快理解账目。以下为产品假设，尚未经过真实店主验证。')
table(['使用机制','本轮状态','为什么可能值得回来'],[
    ['从未核对的账开始','已实现：记录已查看来源与计数；可查看核对摘要。','让用户知道已看到哪里，减少重复翻找；尚无自动“新增账单”检测。'],
    ['关注一件具体的事','已实现：提醒可标记“下次关注”，关注项排前。','把一次发现留在会话中，便于再次核对；尚无处理结果闭环。'],
    ['自己改变到账假设','已实现：0 / 3 / 7 天情景。','店主可判断某笔款慢到时哪个付款日更紧张。'],
    ['带走一份核对摘要','已实现：手动导出账本、测算与提醒 JSON。','用户能留底、复查；尚无 PDF 日结单或一键发送。'],
    ['闭店三分钟核对','设计提案：新增与变化账单摘要、处理完成和下次复查日。','把使用嵌入日常经营节奏，须先解决持久化与新增检测。'],
    ['温和且可控的提醒','设计提案：用户选时间、频率与阈值；同一事项去重。','在付款前提供可用信息，避免重复催促与焦虑。']
],[1.25,2.65,2.6])
h('建议先做的小规模验证')
p('邀请 5—8 位自愿参与的店主，先用虚构或脱敏账单做两轮可用性测试，再考虑两周试用。请店主解释“余额为什么不是销售额”“支出增加来自数量还是单价”“下一笔大额付款前需要核实什么”，记录理解错误与完成时间。')
p('观察首次打开来源并确认的比例、同一周主动返回核对的比例、关注事项被处理的比例、提醒被忽略的原因。目标可以先定为：大多数参与者独立完成首轮核对、能复述净到账公式；这些是验收提议，不是本次成绩。')
p('如果用户只浏览总额、不愿打开来源，就优先缩短核对流程；如果经常抱怨账不全，先改进导入覆盖率；如果提醒无人处理，先调整提醒的时机与可操作性。不要用增加通知数量替代验证。')

page('05  修改过程与工程边界')
table(['阶段','发现与决定','产物'],[
    ['梳理原版','原有核心已有 checkpoint 与 mock 工具；页面展示和实际计算需要接通。','增加版本首页与两个独立界面，复用本地 Node 服务。'],
    ['查资料','结算明细应可追溯；已知账单适合做现金日期表；历史模拟要防止误称预测。','把资料转化为金额口径与产品说明，列出引用 [1]—[3]。'],
    ['修正投资计算','负相关被纳入集中度、费用表与曲线口径不统一、新资产可能漏算。','调整金融工具，增加回归测试；保存完整工具输出。'],
    ['新增店主核心','提醒不能先于账单理解；账本要能去重、解释、重算。','MerchantSession、整数分账本、版本确认、现金表与采购拆解。'],
    ['完善使用路径','一次性报告不足以支持继续核对；信息过多也会掩盖关键数字。','两个产品分别建立重点卡片、可展开详情、关注排序、已查看标记、筛选、导入和手动导出。'],
    ['验证与交付','既要数字正确，也要能打开与操作。','自动化测试、浏览器实测、启动脚本、README、Word 评审文档。']
],[1.05,2.95,2.5])
h('负责的部分与交接点')
p('沿用 Agent Core / Memory 与 Tools / Financial Data 两个责任域：维护任务状态、运行循环、checkpoint、会话记忆、金融与账单计算。新增页面用于展示和验证这些能力，未修改团队仓库其他成员的目录。')
p('投资流程：澄清 → 校验 → 诊断 → 候选方案 → 模拟 → RISK_CHECK 交接。风控决策、审批网关、真实交易仍由对应团队负责。执行 API 固定返回 GATEWAY_NOT_CONNECTED；内部 Paper Trading mock 不当作可信审批。')
h('数据放在哪里')
p('本次仅进程内保存状态。店主按浏览器会话隔离；投资版是本地单用户演示。当前服务仅监听 127.0.0.1，拒绝跨来源 API 请求，不作为公网多用户服务。重启服务会清空账本、审计、记忆与关注状态。')
p('团队仓库保持私有，个人仓库按用户本次确认保持公开。两个仓库均只交付 chenchen/ 的演示代码和说明；不上传原始 PRD、教材、真实账单或登录凭证。')

page('06  验证结果与已知限制')
p('本轮在 Node 环境运行 JavaScript 语法检查与 node --test：31 项自动化测试通过，0 失败。测试用于覆盖关键计算与状态流转，不能推导“没有任何 bug”。')
table(['验证范围','覆盖的关键情形','结果'],[
    ['原有核心：10 项','组合校验、A/B/C、固定 seed、澄清、失败恢复、记忆同意、内部 mock 幂等。','通过'],
    ['投资回归：8 项','错期与伪造字段、负相关、新资产路径、净值费用、旧报告、输入失效、清空偏好、熔断。','通过'],
    ['店主账本：12 项','净额与余额、去重、冲突原子性、日期金额校验、确认门槛、采购拆解、延迟、隔离与空账本。','通过'],
    ['HTTP 集成：1 项','两版页面、Cookie 隔离、确认门槛、失败导入、执行阻断、来源限制、投资主路径与审计。','通过']
],[1.3,4.35,.85])
h('浏览器已经实测的操作')
p('投资版：填写演示偏好并运行，页面展示 31,080 元组合、四组对比和 RISK_CHECK 交接状态；控制台未发现应用错误。店主版：初始确认按钮禁用；筛选外卖账单后打开 2,744 元净额明细；确认后生成提醒；关注状态刷新后仍保留。')
h('上线前还需要补齐')
p('数据：当前投资只有 7 个虚构价格点，店主只有固定示例与手动 JSON。尚未完成真实行情、新闻、银行/支付平台、OCR 或增量流水接入。独立退款、内部转账、混合币种、部分结算和应收坏账等复杂账务仍未支持。')
p('运行：工具调用为同步本地实现，没有远程超时取消、后台调度、持久化重启恢复或多设备同步。商户会话最多保留 500 份，满额后淘汰最早创建的会话。投资的身份鉴别与多租户隔离未实现。')
p('产品：本次是确定性规则原型，未接大模型对话；偏好不驱动个性化风控；提醒需用户打开页面并确认，不自动发送。未上线税务判断、自动付款、账户操作或正式财务报表。')
h('建议的下一轮顺序')
p('先做店主可用性测试与真实账单格式适配；再做数据持久化、可撤销更正和处理闭环；最后在团队明确授权与认证方式后接数据平台、受控通知和风控网关。每步以可回溯账单和可复算结果为验收条件。')

page('07  打开方式、交付位置与参考资料')
h('本地启动')
p('安装 Node.js 22 或更新版本。下载仓库，进入 chenchen，双击“启动两版.cmd”；或在终端运行 npm start。无需 API Key 或第三方运行依赖。保持终端开启，在浏览器访问以下地址：')
table(['页面','地址'],[
    ['版本选择','http://127.0.0.1:4175/'],
    ['持仓健康管理','http://127.0.0.1:4175/investment'],
    ['店主财务助手','http://127.0.0.1:4175/merchant']
],[1.6,4.9])
p('GitHub 用于阅读代码、下载文档与取得可运行项目；文件预览不会运行 Node 后端。其他成员需在自己的电脑启动。若端口占用，可先确认已有服务，或设置 PORT=4180 后启动并访问新端口。')
h('交付位置')
link('团队仓库 · Fintechathon-agent / chenchen（私有）','https://github.com/Fintechathon-agent/Fintechathon-agent/tree/main/chenchen')
link('个人仓库 · realchenchenluo / agent / chenchen（公开）','https://github.com/realchenchenluo/agent/tree/main/chenchen')
p('入口与说明：chenchen/README.md。新页面：versions/。店主计算：merchant/ledger.js。投资工具：tools/financial-tools.js。任务核心：core/harness.js。测试：tests/。本文生成源：scripts/build_report.py。')
h('参考资料：事实与设计推断分开')
link('[1] Stripe · Payout reconciliation report','https://docs.stripe.com/reports/payout-reconciliation')
p('支持“把银行结算款追溯到交易批次、费用和退款”的产品机制。本文据此设计净额拆解，不据此推断国内平台结算规则。')
link('[2] Xero · Small business analytics','https://www.xero.com/us/accounting-software/analytics/')
p('介绍基于已有发票与账单观察未来现金。本文将其简化为 14 天已知账单测算；并未实现其完整分析或自动同步能力。')
link('[3] scikit-learn · Common pitfalls / Data leakage','https://scikit-learn.org/stable/common_pitfalls.html')
p('说明使用预测时不可获得的信息会使评估过于乐观。本文借此区分事后演示比较与样本外验证；并未声称使用了机器学习模型。')
p('资料查阅日期：2026-10-04。持续使用机制、实验建议与功能取舍为本次产品设计推断；来源页面不提供本项目留存效果证据。')

OUT.parent.mkdir(parents=True,exist_ok=True)
for root in [doc.element,doc.styles.element]:
    for border in root.xpath('.//w:pBdr'):
        border.getparent().remove(border)
doc.save(OUT)
print(OUT)
